"""
generate_invoice.py
Generates a sample service invoice Word document (invoice.docx).

Usage:
    python generate_invoice.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import datetime


def set_cell_bg(cell, hex_color: str):
    """Set a table cell background colour using low-level XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def build_invoice(
    output_path: str = "invoice.docx",
    hourly_rate: float = 60.00,
    tax_rate: float = 0.0,
    line_items: list | None = None,
):
    if line_items is None:
        line_items = [
            ("Day 1 – Professional Services", 9.0),
            ("Day 2 – Professional Services", 6.5),
        ]

    doc = Document()

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Header: Company name + INVOICE label ─────────────────────────────────
    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.columns[0].width = Inches(3.5)
    header_table.columns[1].width = Inches(3.0)

    # Left cell – company info
    left = header_table.cell(0, 0)
    left.paragraphs[0].clear()
    company_run = left.paragraphs[0].add_run("Your Company Name")
    company_run.font.size = Pt(18)
    company_run.font.bold = True
    company_run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    for line in ["123 Business Ave", "City, State 00000", "phone: (555) 000-0000", "email: billing@company.com"]:
        p = left.add_paragraph(line)
        p.runs[0].font.size = Pt(9)

    # Right cell – "INVOICE" label
    right = header_table.cell(0, 1)
    right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    invoice_label = right.paragraphs[0]
    invoice_label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = invoice_label.add_run("INVOICE")
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph()  # spacer

    # ── Invoice metadata ──────────────────────────────────────────────────────
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.autofit = False
    meta_table.columns[0].width = Inches(3.25)
    meta_table.columns[1].width = Inches(3.25)

    # Bill To
    bill_cell = meta_table.cell(0, 0)
    bill_cell.paragraphs[0].clear()
    bt_run = bill_cell.paragraphs[0].add_run("Bill To:")
    bt_run.font.bold = True
    bt_run.font.size = Pt(10)
    for line in ["Client Name", "Client Company", "Client Address", "City, State 00000"]:
        p = bill_cell.add_paragraph(line)
        p.runs[0].font.size = Pt(9)

    # Invoice details
    detail_cell = meta_table.cell(0, 1)
    detail_cell.paragraphs[0].clear()
    details = [
        ("Invoice #:", "INV-2024-001"),
        ("Invoice Date:", datetime.date.today().strftime("%B %d, %Y")),
        ("Due Date:", (datetime.date.today() + datetime.timedelta(days=30)).strftime("%B %d, %Y")),
        ("Hourly Rate:", f"${hourly_rate:.2f}"),
    ]
    for label, value in details:
        p = detail_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lbl = p.add_run(f"{label} ")
        lbl.font.bold = True
        lbl.font.size = Pt(9)
        val = p.add_run(value)
        val.font.size = Pt(9)

    doc.add_paragraph()  # spacer

    # ── Line-items table ──────────────────────────────────────────────────────
    HEADER_COLOR = "1F497D"   # dark blue
    ROW_ALT_COLOR = "DCE6F1"  # light blue
    WHITE = "FFFFFF"

    col_widths = [Inches(2.8), Inches(0.9), Inches(1.2), Inches(1.6)]
    headers = ["Description", "Hours", "Rate", "Amount"]

    items_table = doc.add_table(rows=1 + len(line_items), cols=4)
    items_table.style = "Table Grid"

    # Header row
    for col_idx, (heading, width) in enumerate(zip(headers, col_widths)):
        cell = items_table.cell(0, col_idx)
        cell.width = width
        set_cell_bg(cell, HEADER_COLOR)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(heading)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)

    # Data rows
    subtotal = 0.0
    for row_idx, (description, hours) in enumerate(line_items, start=1):
        amount = hours * hourly_rate
        subtotal += amount
        bg = ROW_ALT_COLOR if row_idx % 2 == 0 else WHITE

        values = [description, f"{hours:.1f}", f"${hourly_rate:.2f}", f"${amount:.2f}"]
        aligns = [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
        ]

        for col_idx, (val, align) in enumerate(zip(values, aligns)):
            cell = items_table.cell(row_idx, col_idx)
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            p.alignment = align
            run = p.add_run(val)
            run.font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ── Totals section ────────────────────────────────────────────────────────
    total_hours = sum(h for _, h in line_items)
    tax_amount = subtotal * tax_rate
    total_due = subtotal + tax_amount

    totals_table = doc.add_table(rows=3, cols=2)
    totals_table.autofit = False
    totals_table.columns[0].width = Inches(5.0)
    totals_table.columns[1].width = Inches(1.5)

    totals_data = [
        (f"Subtotal  ({total_hours:.1f} hrs × ${hourly_rate:.2f})", f"${subtotal:.2f}"),
        (f"Tax ({tax_rate:.0%})", f"${tax_amount:.2f}"),
        ("TOTAL DUE", f"${total_due:.2f}"),
    ]

    for row_idx, (label, value) in enumerate(totals_data):
        lbl_cell = totals_table.cell(row_idx, 0)
        val_cell = totals_table.cell(row_idx, 1)

        is_total = row_idx == 2
        bg = "1F497D" if is_total else WHITE
        fc = "FFFFFF" if is_total else "000000"
        set_cell_bg(lbl_cell, bg)
        set_cell_bg(val_cell, bg)

        lp = lbl_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lr = lp.add_run(label)
        lr.font.bold = is_total
        lr.font.size = Pt(10)
        lr.font.color.rgb = RGBColor(*bytes.fromhex(fc))

        vp = val_cell.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        vr = vp.add_run(value)
        vr.font.bold = is_total
        vr.font.size = Pt(10)
        vr.font.color.rgb = RGBColor(*bytes.fromhex(fc))

    doc.add_paragraph()  # spacer

    # ── Notes / Payment instructions ──────────────────────────────────────────
    notes_heading = doc.add_paragraph("Payment Instructions")
    notes_heading.runs[0].font.bold = True
    notes_heading.runs[0].font.size = Pt(10)

    notes = doc.add_paragraph(
        "Please make payment within 30 days of the invoice date.\n"
        "Bank transfer or cheque made payable to 'Your Company Name'.\n"
        "Thank you for your business!"
    )
    notes.runs[0].font.size = Pt(9)

    # ── Footer ────────────────────────────────────────────────────────────────
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("Your Company Name  |  123 Business Ave, City, State  |  (555) 000-0000")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(output_path)
    print(f"Invoice saved to: {output_path}")


if __name__ == "__main__":
    build_invoice("invoice.docx")
